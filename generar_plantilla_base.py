#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Script para generar la plantilla_base.xlsx
Esta plantilla se utiliza para restaurar el archivo de almacén a un estado limpio.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def crear_plantilla_base():
    """
    Crea la plantilla base para el archivo de almacén.
    Incluye las hojas Productos, Movimientos y Reportes con encabezados correctos.
    """
    # Crear nuevo workbook
    wb = Workbook()
    
    # Definir estilos para encabezados
    header_fill = PatternFill(start_color="1b5e20", end_color="1b5e20", fill_type="solid")  # Verde oscuro
    header_font = Font(bold=True, color="FFFFFF")  # Blanco
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # ============================================
    # HOJA 1: PRODUCTOS
    # ============================================
    ws_productos = wb.active
    ws_productos.title = "Productos"
    
    # Definir encabezados de Productos
    headers_productos = ["ID", "Nombre", "Descripción", "Stock", "Unidad de Medida", "Fecha de Creación", "URL Imagen"]
    ws_productos.append(headers_productos)
    
    # Aplicar estilos a los encabezados
    for cell in ws_productos[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border
    
    # Ajustar ancho de columnas
    ws_productos.column_dimensions['A'].width = 8
    ws_productos.column_dimensions['B'].width = 25
    ws_productos.column_dimensions['C'].width = 30
    ws_productos.column_dimensions['D'].width = 12
    ws_productos.column_dimensions['E'].width = 18
    ws_productos.column_dimensions['F'].width = 18
    ws_productos.column_dimensions['G'].width = 30
    
    # ============================================
    # HOJA 2: MOVIMIENTOS
    # ============================================
    ws_movimientos = wb.create_sheet("Movimientos")
    
    # Definir encabezados de Movimientos
    headers_movimientos = [
        "ID", 
        "ID Producto", 
        "Nombre Producto",
        "Tipo", 
        "Cantidad", 
        "Stock Anterior", 
        "Stock Nuevo", 
        "Motivo", 
        "Responsable", 
        "Fecha"
    ]
    ws_movimientos.append(headers_movimientos)
    
    # Aplicar estilos a los encabezados
    for cell in ws_movimientos[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border
    
    # Ajustar ancho de columnas
    ws_movimientos.column_dimensions['A'].width = 8
    ws_movimientos.column_dimensions['B'].width = 12
    ws_movimientos.column_dimensions['C'].width = 25
    ws_movimientos.column_dimensions['D'].width = 12
    ws_movimientos.column_dimensions['E'].width = 12
    ws_movimientos.column_dimensions['F'].width = 15
    ws_movimientos.column_dimensions['G'].width = 15
    ws_movimientos.column_dimensions['H'].width = 20
    ws_movimientos.column_dimensions['I'].width = 20
    ws_movimientos.column_dimensions['J'].width = 15
    
    # ============================================
    # HOJA 3: REPORTES
    # ============================================
    ws_reportes = wb.create_sheet("Reportes")
    
    # Definir encabezados de Reportes
    headers_reportes = ["Fecha de Reporte", "Total Productos", "Total Movimientos", "Stock Total Valor"]
    ws_reportes.append(headers_reportes)
    
    # Aplicar estilos a los encabezados
    for cell in ws_reportes[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border
    
    # Ajustar ancho de columnas
    ws_reportes.column_dimensions['A'].width = 20
    ws_reportes.column_dimensions['B'].width = 20
    ws_reportes.column_dimensions['C'].width = 20
    ws_reportes.column_dimensions['D'].width = 20
    
    # Guardar el archivo
    ruta_salida = os.path.join(os.path.dirname(__file__), "plantilla_base.xlsx")
    wb.save(ruta_salida)
    
    print(f"✅ Plantilla base creada exitosamente en: {ruta_salida}")
    return ruta_salida


if __name__ == "__main__":
    crear_plantilla_base()
