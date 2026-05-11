#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Script para limpiar los datos de prueba del archivo almacen.xlsx
Mantiene la estructura y encabezados, eliminando solo los datos operativos.
"""

import os
from openpyxl import load_workbook


def limpiar_datos_prueba():
    """
    Limpia todos los datos de prueba del archivo almacen.xlsx
    Mantiene encabezados y estructura, dejando el archivo como nuevo.
    """
    ruta_excel = os.path.join(os.path.dirname(__file__), "data", "almacen.xlsx")
    
    if not os.path.exists(ruta_excel):
        print(f"❌ El archivo {ruta_excel} no existe")
        return False
    
    try:
        print(f"📂 Abriendo archivo: {ruta_excel}")
        wb = load_workbook(ruta_excel)
        
        # Limpiar hoja Productos (mantener solo encabezados)
        if "Productos" in wb.sheetnames:
            ws_productos = wb["Productos"]
            # Eliminar todas las filas excepto la de encabezados (fila 1)
            while ws_productos.max_row > 1:
                ws_productos.delete_rows(2)  # Siempre eliminar la fila 2 hasta que solo quede 1
            print("✅ Hoja 'Productos' limpiada (solo encabezados mantenidos)")
        
        # Limpiar hoja Movimientos (mantener solo encabezados)
        if "Movimientos" in wb.sheetnames:
            ws_movimientos = wb["Movimientos"]
            # Eliminar todas las filas excepto la de encabezados (fila 1)
            while ws_movimientos.max_row > 1:
                ws_movimientos.delete_rows(2)  # Siempre eliminar la fila 2 hasta que solo quede 1
            print("✅ Hoja 'Movimientos' limpiada (solo encabezados mantenidos)")
        
        # Limpiar hoja Reportes (mantener solo encabezados)
        if "Reportes" in wb.sheetnames:
            ws_reportes = wb["Reportes"]
            # Eliminar todas las filas excepto la de encabezados (fila 1)
            while ws_reportes.max_row > 1:
                ws_reportes.delete_rows(2)  # Siempre eliminar la fila 2 hasta que solo quede 1
            print("✅ Hoja 'Reportes' limpiada (solo encabezados mantenidos)")
        
        # Guardar el archivo
        wb.save(ruta_excel)
        print(f"\n✅ Archivo limpiado exitosamente: {ruta_excel}")
        print("📝 El archivo está listo para usar como nuevo")
        return True
        
    except Exception as e:
        print(f"❌ Error al limpiar datos: {e}")
        return False


if __name__ == "__main__":
    limpiar_datos_prueba()
